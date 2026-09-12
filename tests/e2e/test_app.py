import re

from openpyxl import load_workbook
from playwright.sync_api import Page, expect
from shiny.playwright import controller
from shiny.pytest import create_app_fixture
from shiny.run import ShinyAppProc

app = create_app_fixture(
    "fixtures/app.py",
    env={
        "AZ_STORAGE_EP": "https://storage.example.com",
        "AZ_STORAGE_RESULTS": "results",
        "AZ_TABLE_ENDPOINT": "https://table.example.com",
        "CAPACITY_MODEL_VERSION": "dev",
        "FEEDBACK_FORM_URL": "",
        "SHINY_TESTMODE": "1",
        "TABLE_NAME": "catalogue",
    },
)


def test_app_displays_capacity_conversion_interface(
    page: Page,
    app: ShinyAppProc,
) -> None:
    page.goto(app.url)

    expect(page).to_have_title("OpenPlan Capacity Conversion Model")
    favicon = page.locator("head link[rel='icon']")
    expect(favicon).to_have_attribute("href", "favicon.ico")
    favicon_response = page.request.get(f"{app.url}/favicon.ico")
    assert favicon_response.ok
    assert favicon_response.headers["content-type"] in {
        "image/vnd.microsoft.icon",
        "image/x-icon",
    }

    banner = page.get_by_role("banner")
    expect(
        banner.get_by_text("OpenPlan Capacity Conversion Model", exact=True)
    ).to_be_visible()
    expect(page.get_by_role("heading", name="Capacity estimates")).to_be_visible()
    logo = page.get_by_role("img", name="The Strategy Unit and NHS")
    dataset = controller.InputSelect(page, "dataset")
    scenario = controller.InputSelect(page, "scenario")
    model_run = controller.InputSelect(page, "model_run")
    generate = controller.InputActionButton(page, "generate")
    estimates = controller.OutputDataFrame(page, "estimates")
    feedback = controller.InputActionButton(page, "feedback")
    download = controller.DownloadButton(page, "download_estimates")
    results_card = page.locator(".card").filter(
        has=page.locator(".card-header", has_text="Capacity estimates")
    )

    expect(logo).to_be_visible()
    expect(logo).to_have_js_property("complete", True)
    expect(logo).to_have_js_property("naturalWidth", 2000)
    expect(results_card).to_be_visible()
    results_card_box = results_card.bounding_box()
    assert results_card_box is not None
    assert page.viewport_size is not None
    assert results_card_box["width"] >= page.viewport_size["width"] * 0.95
    dataset.expect_choices(["", "RXX"])
    dataset.set("RXX")
    scenario.expect_choices(["", "Example scenario"])
    scenario.set("Example scenario")
    model_run.expect_choices(["", "test-guid"])
    expect(model_run.loc_choices.nth(1)).to_have_text("17 Aug 2026, 14:37 UTC")
    model_run.set("test-guid")
    generate.click()

    expect(estimates.loc).to_be_visible()
    expect(estimates.loc).to_contain_text("ip_wards")
    expect(estimates.loc).to_contain_text("ip_procedures_and_theatres")
    expect(feedback.loc).to_be_visible()
    expect(feedback.loc).to_have_class(re.compile(r"\bbtn-sm\b"))
    expect(download.loc).to_be_visible()
    expect(download.loc).to_have_class(re.compile(r"\bbtn-sm\b"))
    expect(download.loc).to_have_attribute("href", re.compile(r".+"))

    with page.expect_download() as download_info:
        download.click()
    downloaded_workbook = download_info.value
    assert downloaded_workbook.suggested_filename == "capacity_conversion_results.xlsx"
    workbook_path = downloaded_workbook.path()
    assert workbook_path is not None
    with workbook_path.open("rb") as workbook_file:
        workbook = load_workbook(workbook_file, read_only=True)
        assert workbook.sheetnames == [
            "coversheet",
            "metadata",
            "assumptions",
            "baseline_year_activity_counts",
            "predicted_activity_volumes",
            "estimated_capacity_needs",
        ]
        metadata_keys = [
            row[0] for row in workbook["metadata"].iter_rows(values_only=True)
        ]
        assert metadata_keys[:6] == [
            "dataset",
            "capacity_model_version",
            "ip_sites",
            "op_sites",
            "aae_sites",
            "capacity_conversion_runtime",
        ]

    page_count = len(page.context.pages)
    feedback.click()

    feedback_dialog = page.get_by_role("dialog")
    close = feedback_dialog.get_by_role("button", name="Close")
    expect(feedback_dialog).to_be_visible()
    expect(feedback_dialog).to_contain_text(
        "The feedback form is not currently available."
    )
    expect(close).to_have_class(re.compile(r"\bbtn-primary\b"))
    expect(close).to_have_class(re.compile(r"\bbtn-sm\b"))
    assert len(page.context.pages) == page_count
