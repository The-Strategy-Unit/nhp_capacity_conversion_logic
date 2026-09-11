from collections import OrderedDict

import pandas as pd
import pytest
from openpyxl.workbook.workbook import Workbook
from pandas.testing import assert_frame_equal, assert_series_equal

from nhp.capacity_conversion.results import (
    add_care_setting_and_summarise,
    add_coversheet,
    apply_styling_to_coversheet,
    combine_and_format_dataframes,
    create_and_format_baseline_df,
    create_and_format_capacity_needs_df,
    create_and_format_predicted_vols_df,
    process_and_save_results_to_excel,
    process_data_to_save,
    summarise_model_runs,
    tidy_metadata,
)


def test_process_and_save_results_to_excel(mocker):
    # arrange

    mock_makedirs = mocker.patch("nhp.capacity_conversion.results.os.makedirs")
    mocker.patch(
        "nhp.capacity_conversion.results.os.path.join",
        side_effect=lambda *x: "/".join(x),
    )
    mock_wb = mocker.Mock()
    mock_ws = mocker.Mock()
    mock_cell = mocker.Mock()
    mock_cell.value = "val"
    mock_cell.column_letter = "A"
    mock_ws.columns = [(mock_cell,), (mock_cell,)]
    mock_ws.column_dimensions = {"A": mocker.Mock()}

    mocker.patch("nhp.capacity_conversion.results.Workbook", return_value=mock_wb)
    mock_wb.active = mocker.Mock()
    mock_wb.create_sheet.return_value = mock_ws

    mock_dataframe_to_rows = mocker.patch(
        "nhp.capacity_conversion.results.dataframe_to_rows",
        side_effect=[
            [
                ["guid", "123"],
                ["capacity_conversion_runtime", "456"],
            ],
            [
                ["col1", "col2"],
                ["val1", "val2"],
            ],
        ],
    )
    mock_logger = mocker.patch("nhp.capacity_conversion.results.logger")
    metadata = pd.Series(
        {
            "guid": "123",
            "capacity_conversion_runtime": "456",
        }
    )
    df = pd.DataFrame(
        {
            "model_run": list(range(11)),
            "group": ["group"] * 11,
            "value": list(range(11)),
        }
    ).set_index(["model_run", "group"])
    data_to_save = {
        "metadata": metadata,
        "results": df,
    }
    mock_process_data_to_save = mocker.patch(
        "nhp.capacity_conversion.results.process_data_to_save",
        return_value=data_to_save,
    )
    mock_apply_styling_to_coversheet = mocker.patch(
        "nhp.capacity_conversion.results.apply_styling_to_coversheet"
    )

    # act
    process_and_save_results_to_excel(data_to_save)

    # assert
    mock_makedirs.assert_called_once_with("results/123/456", exist_ok=True)
    mock_wb.remove.assert_called_once_with(mock_wb.active)
    assert mock_wb.create_sheet.call_count == len(data_to_save)
    assert mock_dataframe_to_rows.call_count == 2
    mock_process_data_to_save.assert_called_once()

    # metadata is a Series, so it should be written without headers
    metadata_call = mock_dataframe_to_rows.call_args_list[0]
    assert metadata_call.kwargs["index"] is False
    assert metadata_call.kwargs["header"] is False

    # results is a DataFrame, so it should include headers
    results_call = mock_dataframe_to_rows.call_args_list[1]
    assert results_call.kwargs["index"] is False
    assert results_call.kwargs["header"] is True

    mock_apply_styling_to_coversheet.assert_called_once()

    mock_wb.save.assert_called_once_with(
        "results/123/456/capacity_conversion_results.xlsx"
    )
    mock_logger.info.assert_called_once()


def test_summarise_model_runs():
    df = pd.DataFrame(
        {
            "model_run": list(range(11)),
            "group": ["group"] * 11,
            "value": list(range(11)),
        }
    ).set_index(["model_run", "group"])
    expected = pd.DataFrame(
        {"group": ["group"], "p10": [1.0], "mean": [5.0], "p90": [9.0]}
    ).set_index("group")
    actual = summarise_model_runs(df)
    assert_frame_equal(actual, expected)


def test_summarise_model_runs_with_multiple_cols():
    df = pd.DataFrame(
        {
            "model_run": list(range(11)),
            "grouping": ["group"] * 11,
            "value": list(range(11)),
            "value_2": list(range(11)),
        }
    ).set_index(["model_run", "grouping"])
    actual = summarise_model_runs(df)
    assert actual.index.names == ["grouping", "measure"]
    assert list(actual.index.get_level_values("measure").unique()) == [
        "value",
        "value_2",
    ]


def test_summarise_model_runs_with_multiple_indexes():
    df = pd.DataFrame(
        {
            "model_run": list(range(11)),
            "group": ["group"] * 11,
            "value": list(range(11)),
            "index_2": list(range(11)),
        }
    ).set_index(["model_run", "group", "index_2"])
    with pytest.raises(ValueError, match="Expected exactly one index column."):
        summarise_model_runs(df)


def test_tidy_metadata():
    metadata = pd.Series(
        {
            "dataset": "dataset",
            "capacity_model_version": "1.2.3",
            "ip_sites": "A,B",
            "op_sites": "ALL",
            "aae_sites": "ALL",
            "capacity_conversion_runtime": "capacity_conversion_runtime",
            "app_version": "4.5.6",
            "scenario_name": "scenario_name",
            "scenario_runtime": "scenario_runtime",
            "unwanted_metadata": "remove me",
        }
    )
    results = pd.DataFrame({"value": [1, 2, 3]})

    data_to_save = {"metadata": metadata, "results": results}

    result = tidy_metadata(data_to_save)

    expected = pd.Series(
        {
            "dataset": "dataset",
            "capacity_model_version": "1.2.3",
            "ip_sites": "A,B",
            "op_sites": "ALL",
            "aae_sites": "ALL",
            "capacity_conversion_runtime": "capacity_conversion_runtime",
            "demand_model_version": "4.5.6",
            "demand_model_scenario_name": "scenario_name",
            "demand_model_scenario_runtime": "scenario_runtime",
        }
    )

    # Test that metadata has changed
    assert_series_equal(result["metadata"], expected)  # ty:ignore invalid-argument-type
    # Test that results remain unchanged
    assert_frame_equal(result["results"], results)  # ty:ignore invalid-argument-type


def test_tidy_metadata_if_no_metadata():
    results = pd.DataFrame({"value": [1, 2, 3]})
    data_to_save = {"results": results}
    result = tidy_metadata(data_to_save)
    assert_frame_equal(result["results"], results)  # ty:ignore invalid-argument-type
    assert "results" in result


def test_add_coversheet():
    data_to_save = {
        "metadata": pd.Series({"capacity_conversion_runtime": "20260101_111111"}),
        "assumptions": pd.DataFrame({"value": [1, 2, 3]}),
    }

    result = add_coversheet(data_to_save)

    assert isinstance(result, OrderedDict)
    assert list(result.keys()) == ["coversheet", "metadata", "assumptions"]

    coversheet = result["coversheet"]
    assert isinstance(coversheet, pd.Series)
    assert coversheet["Date Created"] == "01/01/2026"
    assert coversheet["OpenPlan Capacity Conversion Results Workbook"] == ""
    assert "OpenPlan capacity conversion model" in coversheet["Introduction"]


def test_apply_styling_to_coversheet():
    workbook = Workbook()
    ws = workbook.active
    ws.title = "coversheet"

    ws["A1"] = "OpenPlan Capacity Conversion Results Workbook"
    ws["A2"] = "Date Created"
    ws["A3"] = "Introduction"
    ws["A4"] = "Contents"
    ws["A5"] = "do_not_bold"

    apply_styling_to_coversheet(workbook)

    assert all(ws[f"A{row}"].font.bold for row in range(1, 5))
    assert ws["A5"].font.bold is False


def test_add_care_setting_and_summarise(mocker):
    input_df = pd.DataFrame(
        {"value": [10, 20]},
        index=pd.MultiIndex.from_tuples(
            [
                (1, "A"),
                (2, "A"),
            ],
            names=["model_run", "activity"],
        ),
    )

    summarised_df = pd.DataFrame(
        {"value": [30]},
        index=pd.Index(["A"], name="activity"),
    )

    input_no_model_run = pd.DataFrame(
        {"value": [10, 20]},
        index=pd.MultiIndex.from_tuples(
            [
                ("A", "B"),
                ("A", "B"),
            ],
            names=["activity", "measure"],
        ),
    )

    mock_summarise = mocker.patch(
        "nhp.capacity_conversion.results.summarise_model_runs",
        return_value=summarised_df,
    )

    data_to_save = {
        "ip_capacity": input_df,
        "ip_baseline": input_no_model_run,
        "metadata": pd.Series({"guid": "test-guid"}),
    }

    result = add_care_setting_and_summarise(data_to_save)

    mock_summarise.assert_called_once_with(input_df)

    expected_capacity = pd.DataFrame(
        {
            "value": [30],
            "care_setting": ["ip"],
        },
        index=pd.Index(["A"], name="activity"),
    )
    output_no_model_run = pd.DataFrame(
        {"value": [10, 20], "care_setting": "ip"},
        index=pd.MultiIndex.from_tuples(
            [
                ("A", "B"),
                ("A", "B"),
            ],
            names=["activity", "measure"],
        ),
    )

    assert_frame_equal(result["ip_capacity"], expected_capacity)  # ty: ignore
    assert_frame_equal(result["ip_baseline"], output_no_model_run)  # ty:ignore
    assert result["metadata"].equals(data_to_save["metadata"])


def test_create_and_format_baseline_df():
    df = pd.DataFrame(
        {
            "care_setting": ["a"],
            "total": [10],
            "spells": [5],
            "beddays": [20],
            "total_theatre_time": [None],
        },
        index=pd.Index(["activity_group_measure"], name="activity_group"),
    )

    result = create_and_format_baseline_df([df])

    expected = pd.DataFrame(
        {
            "care_setting": ["a", "a", "a"],
            "value": [10, 5, 20],
        },
        index=pd.MultiIndex.from_tuples(
            [
                ("activity_group_measure", "measure"),
                ("activity_group_measure", "spells"),
                ("activity_group_measure", "beddays"),
            ],
            names=["activity_group", "measure"],
        ),
    )

    print(result)
    print(expected)
    assert_frame_equal(result, expected, check_dtype=False)


def test_create_and_format_predicted_vols_df():
    df1 = pd.DataFrame(
        {"value": [1, 2]},
        index=pd.MultiIndex.from_tuples(
            [
                ("activity_group_a", "measure_a"),
                ("activity_group_a", "measure_b"),
            ],
            names=["index", "measure"],
        ),
    )
    df2 = pd.DataFrame(
        {"value": [1, 2]},
        index=pd.Index(
            [
                "activity_group_attendances",
                "activity_group_attendances",
            ],
            name="index",
        ),
    )

    result = create_and_format_predicted_vols_df([df1, df2])

    expected = pd.DataFrame(
        {"value": [1, 2] * 2},
        index=pd.MultiIndex.from_tuples(
            [
                ("activity_group_a", "measure_a"),
                ("activity_group_a", "measure_b"),
                ("activity_group_attendances", "attendances"),
                ("activity_group_attendances", "attendances"),
            ],
            names=["activity_group", "measure"],
        ),
    )

    assert_frame_equal(result, expected)


def test_create_and_format_capacity_needs_df():
    df1 = pd.DataFrame({"value": [1.5]}, index=pd.Index(["a"]))
    df2 = pd.DataFrame({"value": [2.5]}, index=pd.Index(["b"]))

    result = create_and_format_capacity_needs_df([df1, df2])

    expected = pd.DataFrame(
        {"value": [1.5, 2.5]},
        index=pd.Index(["a", "b"], name="resource"),
    )

    assert_frame_equal(result, expected)


def test_combine_and_format_dataframes(mocker):
    baseline_result = pd.DataFrame({"value": [1, 1]})
    predicted_result = pd.DataFrame({"value": [2, 2]})
    capacity_result = pd.DataFrame({"value": [3, 3]})

    mock_baseline = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_baseline_df",
        return_value=baseline_result,
    )
    mock_predicted = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_predicted_vols_df",
        return_value=predicted_result,
    )
    mock_capacity = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_capacity_needs_df",
        return_value=capacity_result,
    )

    data = {
        "a_baseline": pd.DataFrame({"value": [1]}),
        "b_baseline": pd.DataFrame({"value": [1]}),
        "a_fun_area_groupings": pd.DataFrame({"value": [2]}),
        "b_fun_area_groupings": pd.DataFrame({"value": [2]}),
        "a_capacity": pd.DataFrame({"value": [3]}),
        "b_capacity": pd.DataFrame({"value": [3]}),
        "metadata": pd.DataFrame({"value": [4]}),
    }

    result = combine_and_format_dataframes(data)

    mock_baseline.assert_called_once()
    mock_predicted.assert_called_once()
    mock_capacity.assert_called_once()

    assert_frame_equal(result["baseline_year_activity_counts"], baseline_result)  # ty: ignore
    assert_frame_equal(
        result["predicted_activity_volumes"],  # ty: ignore
        predicted_result.round(2),
    )
    assert_frame_equal(
        result["estimated_capacity_needs"],  # ty: ignore
        capacity_result.round(2),
    )
    assert set(result) == {
        "baseline_year_activity_counts",
        "predicted_activity_volumes",
        "estimated_capacity_needs",
        "metadata",
    }


def test_combine_and_format_dataframes_no_keys(mocker):

    mock_baseline = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_baseline_df",
    )
    mock_predicted = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_predicted_vols_df",
    )
    mock_capacity = mocker.patch(
        "nhp.capacity_conversion.results.create_and_format_capacity_needs_df",
    )

    data = {
        "metadata": pd.DataFrame({"value": [4]}),
    }

    result = combine_and_format_dataframes(data)

    mock_baseline.assert_not_called()
    mock_predicted.assert_not_called()
    mock_capacity.assert_not_called()

    assert set(result) == {
        "metadata",
    }


def test_process_data_to_save(mocker):
    data_to_save = {"data_to_save": pd.DataFrame()}
    mock_add_care_setting_and_summarise = mocker.patch(
        "nhp.capacity_conversion.results.add_care_setting_and_summarise",
        return_value=data_to_save,
    )
    mock_combine_and_format_dataframes = mocker.patch(
        "nhp.capacity_conversion.results.combine_and_format_dataframes",
        return_value=data_to_save,
    )
    mock_tidy_metadata = mocker.patch(
        "nhp.capacity_conversion.results.tidy_metadata",
        return_value=data_to_save,
    )
    mock_add_coversheet = mocker.patch(
        "nhp.capacity_conversion.results.add_coversheet",
        return_value=data_to_save,
    )

    process_data_to_save(data_to_save)

    mock_add_care_setting_and_summarise.assert_called_once()
    mock_combine_and_format_dataframes.assert_called_once()
    mock_tidy_metadata.assert_called_once()
    mock_add_coversheet.assert_called_once()
